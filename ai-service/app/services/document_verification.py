"""Service de vérification automatique de documents d'onboarding.

Analyse les documents de constitution du dossier (``CV``, ``CIN``,
``DEMANDE_STAGE``, ``LETTRE_AFFECTATION``) et produit un résultat de
vérification structuré, sans jamais affirmer l'authenticité du document :

- extraction du texte lisible (PDF via ``pypdf``, DOCX via ``python-docx``,
  DOC via ``textutil`` si disponible, images via OCR optionnel) ;
- identification du type de document (mots-clés + indice du nom de fichier) ;
- vérification de cohérence avec les informations du candidat fournies par le
  backend Node (nom/prénom, numéro CIN le cas échéant) ;
- présence des informations requises pour le type attendu ;
- score de confiance (0-100, plafonné à 95 : on ne certifie jamais un
  document à 100 %) et statut ``VALID`` / ``REVIEW_REQUIRED`` / ``REJECTED``.

Le résultat reste consultable par un RH : **aucune décision finale n'est
automatisée** sur la seule foi de cette vérification.

Sécurité : le numéro CIN, le contenu complet des documents, les mots de passe
et les jetons ne sont JAMAIS journalisés. Seuls ``candidature_id``, le type
attendu/détecté, le statut, la confiance et le nombre d'anomalies sont loggés.

Erreurs normalisées :

- ``DocumentValidationError``   -> HTTP 400 (fichier vide, trop volumineux,
  extension non prise en charge, type de document inconnu) ;
- ``DocumentUnavailableError``  -> HTTP 503 (OCR requis mais indisponible).
"""

import io
import logging
import os
import re
import shutil
import subprocess
import unicodedata
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

from ..schemas.documents import (
    CandidateInfo,
    DocumentChecks,
    DocumentVerificationResponse,
    VerdictStatus,
)

# ─────────────────────────────── Contrats ────────────────────────────────────

# Types de documents d'onboarding acceptés (identiques au backend Node :
# backend/validators/application.validator.js).
DOCUMENT_TYPES = ("CV", "CIN", "DEMANDE_STAGE", "LETTRE_AFFECTATION")

# Libellés français pour les messages (sans donnée personnelle).
DOCUMENT_LABELS = {
    "CV": "un curriculum vitae",
    "CIN": "une carte nationale d'identité",
    "DEMANDE_STAGE": "une demande de stage",
    "LETTRE_AFFECTATION": "une lettre d'affectation",
}

# Extensions de fichiers prises en charge (alignées sur le backend Node).
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".png", ".jpg", ".jpeg", ".txt"}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}

# Nombre de chiffres d'une carte d'identité tunisienne.
CIN_PATTERN = re.compile(r"\b\d{8}\b")

# ────────────────────────────────── Erreurs ──────────────────────────────────


class DocumentValidationError(Exception):
    """Fichier invalide (vide, trop volumineux, format non pris en charge)."""


class DocumentUnavailableError(Exception):
    """Le service de vérification est indisponible (ex. OCR requis et absent)."""


# ─────────────────────── Extraction du texte lisible ─────────────────────────

@dataclass(frozen=True)
class ExtractedText:
    """Texte extrait d'un document + nature du support analysé."""

    text: str
    kind: str  # "pdf", "docx", "doc", "txt", "image"


def _extract_text(content: bytes, ext: str) -> ExtractedText:
    """Extrait le texte lisible du document (délégation par format)."""
    if ext in IMAGE_EXTENSIONS:
        return ExtractedText(text=_ocr_image(content), kind="image")
    if ext == ".pdf":
        return ExtractedText(text=_extract_pdf(content), kind="pdf")
    if ext == ".docx":
        return ExtractedText(text=_extract_docx(content), kind="docx")
    if ext == ".doc":
        return ExtractedText(text=_extract_doc(content), kind="doc")
    if ext == ".txt":
        return ExtractedText(text=_decode_text(content), kind="txt")
    return ExtractedText(text="", kind=ext)


def _extract_pdf(content: bytes) -> str:
    """Texte extrait d'un PDF via pypdf ('' si le PDF est corrompu/illisible)."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        parts: List[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    except Exception as exc:  # PdfReadError, ValueError… => document illisible.
        logger.warning("[DOC-VERIFY] extraction PDF impossible : %s", type(exc).__name__)
        return ""


def _extract_docx(content: bytes) -> str:
    """Texte extrait d'un DOCX via python-docx (paragraphes + tableaux)."""
    try:
        import docx

        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as exc:  # document corrompu / non reconnu.
        logger.warning("[DOC-VERIFY] extraction DOCX impossible : %s", type(exc).__name__)
        return ""


def _extract_doc(content: bytes) -> str:
    """Texte d'un DOC via ``textutil`` (macOS). Vide si l'outil est absent."""
    textutil = shutil.which("textutil")
    if not textutil:
        logger.warning("[DOC-VERIFY] extraction DOC impossible : textutil indisponible.")
        return ""
    try:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".doc") as src, tempfile.NamedTemporaryFile(
            suffix=".txt", mode="r"
        ) as dst:
            src.write(content)
            src.flush()
            subprocess.run(
                [textutil, "-convert", "txt", "-output", dst.name, src.name],
                check=True,
                capture_output=True,
                timeout=30,
            )
            dst.seek(0)
            return dst.read()
    except (OSError, subprocess.SubprocessError) as exc:
        logger.warning("[DOC-VERIFY] extraction DOC impossible : %s", exc)
        return ""


def _decode_text(content: bytes) -> str:
    """Décode un fichier texte simple (UTF-8 puis latin-1)."""
    for encoding in ("utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _ocr_image(content: bytes) -> str:
    """OCR d'une image via tesseract (optionnel). '' si OCR indisponible."""
    try:
        from PIL import Image  # noqa: F401
        import pytesseract
    except ImportError:
        logger.info("[DOC-VERIFY] OCR indisponible (pytesseract/Pillow absent).")
        return ""
    if not shutil.which("tesseract"):
        logger.info("[DOC-VERIFY] OCR indisponible (binaire tesseract absent).")
        return ""
    try:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".png") as tmp:
            tmp.write(content)
            tmp.flush()
            return pytesseract.image_to_string(tmp.name, lang="fra").strip()
    except Exception as exc:  # image corrompue, langue manquante…
        logger.warning("[DOC-VERIFY] OCR échoué : %s", type(exc).__name__)
        return ""


def _ocr_available() -> bool:
    """Un moteur OCR (tesseract) est-il réellement utilisable ?"""
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False
    return bool(shutil.which("tesseract"))


# ──────────────────── Identification du type de document ─────────────────────

# Mots-clés distinctifs par type (poids fort pour les expressions précises).
TYPE_KEYWORDS: Dict[str, Dict[str, int]] = {
    "CV": {
        "curriculum vitae": 5,
        "curriculum": 3,
        "experience professionnelle": 4,
        "expérience professionnelle": 4,
        "formation": 2,
        "competences": 3,
        "compétences": 3,
        "parcours": 2,
        "cv": 4,
    },
    "CIN": {
        "carte nationale d'identite": 6,
        "carte nationale d'identité": 6,
        "carte d'identite": 6,
        "carte d'identité": 6,
        "numero de carte": 5,
        "numéro de carte": 5,
        "carte d identite": 6,
        "carte nationale d identite": 6,
        "cin": 3,
    },
    "DEMANDE_STAGE": {
        "demande de stage": 6,
        "demande d'un stage": 6,
        "demande d un stage": 6,
        "demande d'admission": 4,
        "demande d admission": 4,
        "candidature à un stage": 5,
        "candidature a un stage": 5,
        "stage de perfectionnement": 4,
    },
    "LETTRE_AFFECTATION": {
        "lettre d'affectation": 6,
        "lettre d affectation": 6,
        "affectation": 5,
        "direction des ressources humaines": 5,
        "ressources humaines": 3,
        "steg": 2,
    },
}

# Chaque carte d'identité porte un numéro à 8 chiffres : indice supplémentaire.
CIN_NUMBER_HINT_WEIGHT = 2


def _normalize(text: str) -> str:
    """Minuscules + suppression des accents pour des comparaisons robustes."""
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(c for c in text if not unicodedata.combining(c))
    return text.lower()


def _filename_hints(filename: Optional[str]) -> Dict[str, int]:
    """Indices de type déduits du nom de fichier (ex. ``cv-salma.pdf``)."""
    base = os.path.splitext(filename or "")[0].lower()
    hints: Dict[str, int] = {}
    if "curriculum" in base or base == "cv" or "cv" in base.replace("-", " ").replace("_", " ").split():
        hints["CV"] = hints.get("CV", 0) + 4
    if any(token in base for token in ("cin", "cni", "carte", "identite", "identité")):
        hints["CIN"] = hints.get("CIN", 0) + 4
    if "demande" in base or "motivation" in base:
        hints["DEMANDE_STAGE"] = hints.get("DEMANDE_STAGE", 0) + 4
    if "affectation" in base or "lettre" in base:
        hints["LETTRE_AFFECTATION"] = hints.get("LETTRE_AFFECTATION", 0) + 4
    return hints


def _has_keyword(text_norm: str, keyword: str) -> bool:
    """Un mot-clé apparaît-il dans le texte normalisé ?

    Les expressions (avec espaces) sont cherchées en sous-chaîne ; les mots
    simples (« cv », « cin »…) le sont avec des limites de mot pour éviter les
    faux positifs à l'intérieur d'autres mots.
    """
    if " " in keyword:
        return keyword in text_norm
    return bool(re.search(rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])", text_norm))


def _detect_document_type(text_norm: str, filename: Optional[str]) -> Optional[str]:
    """Type de document le plus probable, ou ``None`` si non identifiable.

    Score = somme pondérée des mots-clés présents dans le texte normalisé,
    complétée par les indices du nom de fichier et par la présence d'un
    numéro à 8 chiffres (fort indice de carte d'identité).
    """
    scores: Dict[str, int] = {t: 0 for t in DOCUMENT_TYPES}

    for doc_type, keywords in TYPE_KEYWORDS.items():
        for keyword, weight in keywords.items():
            if _has_keyword(text_norm, keyword):
                scores[doc_type] += weight

    if CIN_PATTERN.search(text_norm):
        scores["CIN"] += CIN_NUMBER_HINT_WEIGHT

    for doc_type, weight in _filename_hints(filename).items():
        scores[doc_type] += weight

    best_type, best_score = max(scores.items(), key=lambda item: item[1])
    if best_score <= 0:
        return None
    return best_type


# ──────────────────── Cohérence avec les informations candidat ───────────────

def _contains_token(text_norm: str, token: str) -> bool:
    """Un mot (ou une expression) apparaît-il dans le texte normalisé ?"""
    token = _normalize(token).strip()
    if not token:
        return False
    if " " in token:
        return token in text_norm
    return bool(re.search(rf"(?<![a-z0-9]){re.escape(token)}(?![a-z0-9])", text_norm))


def _match_candidate_name(text_norm: str, candidate: CandidateInfo) -> Tuple[str, Optional[str]]:
    """Évalue la présence du nom/prénom du candidat dans le document.

    Renvoie ``("full", …)``, ``("partial", …)``, ``("none", …)`` ou
    ``("unverifiable", None)`` si aucune identité n'a été fournie.
    """
    if not (candidate.nom or candidate.prenom):
        return "unverifiable", None

    surname_found = bool(candidate.nom) and _contains_token(text_norm, candidate.nom)
    first_found = bool(candidate.prenom) and _contains_token(text_norm, candidate.prenom)

    if candidate.nom and candidate.prenom:
        if surname_found and first_found:
            return "full", None
        if surname_found or first_found:
            return "partial", "Seule une partie du nom du candidat a été retrouvée dans le document."
        return "none", "Le nom du candidat ne figure pas dans le document."
    if surname_found or first_found:
        return "full", None
    return "none", "Le nom du candidat ne figure pas dans le document."


def _find_cin_numbers(text: str) -> List[str]:
    """Numéros candidats à 8 chiffres présents dans le texte."""
    return CIN_PATTERN.findall(text)


def _check_cin(cin_numbers: List[str], candidate: CandidateInfo) -> Tuple[bool, Optional[str]]:
    """Compare le numéro trouvé avec celui du candidat (documents CIN uniquement)."""
    expected = re.sub(r"\D", "", candidate.cin or "")
    if not expected:
        return True, None  # aucune référence à comparer
    if not cin_numbers:
        return False, "Le numéro de la carte d'identité est introuvable dans le document."
    if any(number == expected for number in cin_numbers):
        return True, None
    return False, "Le numéro de la carte d'identité ne correspond pas aux informations du candidat."


# ──────────────── Présence des informations requises par type ────────────────

# Pour le CV, la « formation » (diplômes) ET l'« expérience » (stages, projets)
# sont attendues. Pour la CIN, un numéro à 8 chiffres. Pour la demande de
# stage, la mention « demande ». Pour la lettre d'affectation, « affectation ».
CV_FORMATION_KEYWORDS = (
    "formation", "diplome", "diplôme", "licence", "master", "ingenieur",
    "ingénieur", "universite", "université", "ecole", "école", "bac", "études",
)
CV_EXPERIENCE_KEYWORDS = (
    "experience", "expérience", "stage", "projet", "competence", "compétence",
    "langage", "connaissance",
)


def _check_required_information(
    expected_type: str, text_norm: str, cin_numbers: List[str]
) -> Tuple[bool, Optional[str]]:
    """Les informations obligatoires pour le type attendu sont-elles présentes ?"""
    if expected_type == "CV":
        has_formation = any(k in text_norm for k in CV_FORMATION_KEYWORDS)
        has_experience = any(k in text_norm for k in CV_EXPERIENCE_KEYWORDS)
        if has_formation and has_experience:
            return True, None
        return False, (
            "Informations requises manquantes dans le CV : parcours de formation "
            "et expérience (stages, projets, compétences) attendus."
        )
    if expected_type == "CIN":
        if cin_numbers:
            return True, None
        return False, "Informations requises manquantes : aucun numéro de carte d'identité lisible."
    if expected_type == "DEMANDE_STAGE":
        if "demande" in text_norm or ("objet" in text_norm and "stage" in text_norm):
            return True, None
        return False, "Informations requises manquantes : la mention « demande de stage » est absente."
    if expected_type == "LETTRE_AFFECTATION":
        if "affectation" in text_norm:
            return True, None
        return False, "Informations requises manquantes : la mention « affectation » est absente."
    return True, None


# ──────────────────────────────── Service ────────────────────────────────────


class DocumentVerificationService:
    """Point d'entrée de la vérification de documents.

    ``verify_document`` reçoit le fichier, le type attendu et les informations
    du candidat ; il renvoie un ``DocumentVerificationResponse`` structuré.
    """

    def __init__(self, max_bytes: int, ocr_required: bool = False):
        self._max_bytes = max_bytes
        self._ocr_required = ocr_required
        if ocr_required and not _ocr_available():
            raise DocumentUnavailableError(
                "Le service de vérification requiert l'OCR (DOCUMENT_OCR_REQUIRED) "
                "mais tesseract n'est pas disponible."
            )

    def _validate_input(self, content: bytes, filename: str, expected_type: str) -> None:
        if not content:
            raise DocumentValidationError("Le fichier est vide.")
        if len(content) > self._max_bytes:
            raise DocumentValidationError(
                f"Fichier trop volumineux ({len(content) / 1024 / 1024:.1f} Mo, "
                f"max {self._max_bytes // 1024 // 1024} Mo)."
            )
        ext = os.path.splitext(filename or "")[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise DocumentValidationError(
                f"Format non pris en charge ({filename or ext}). Formats acceptés : "
                + ", ".join(sorted(SUPPORTED_EXTENSIONS))
                + "."
            )
        if expected_type not in DOCUMENT_TYPES:
            raise DocumentValidationError(
                f"Type de document inconnu ({expected_type}). Types acceptés : "
                + ", ".join(DOCUMENT_TYPES)
                + "."
            )

    def verify_document(
        self,
        content: bytes,
        filename: Optional[str],
        expected_type: str,
        candidate: Optional[CandidateInfo] = None,
        candidature_id: str = "",
    ) -> DocumentVerificationResponse:
        """Vérifie un document et renvoie le résultat structuré."""
        candidate = candidate or CandidateInfo()
        self._validate_input(content, filename or "", expected_type)

        ext = os.path.splitext(filename or "")[1].lower()
        extracted = _extract_text(content, ext)
        text_norm = _normalize(extracted.text)

        detected_type = _detect_document_type(text_norm, filename)
        name_status, name_issue = _match_candidate_name(text_norm, candidate)
        cin_numbers = _find_cin_numbers(extracted.text)
        cin_ok, cin_issue = _check_cin(cin_numbers, candidate)
        required_ok, required_issue = _check_required_information(expected_type, text_norm, cin_numbers)

        status, confidence, issues = self._decide(
            detected_type=detected_type,
            expected_type=expected_type,
            extracted=extracted,
            name_status=name_status,
            required_ok=required_ok,
            cin_ok=cin_ok,
            cin_issue=cin_issue,
            name_issue=name_issue,
            required_issue=required_issue,
        )

        checks = DocumentChecks(
            document_type_match=detected_type == expected_type,
            candidate_name_match=name_status in ("full", "unverifiable"),
            required_information_present=required_ok,
        )

        logger.info(
            "[DOC-VERIFY] candidature=%s expected=%s detected=%s status=%s confidence=%d issues=%d",
            candidature_id or "-",
            expected_type,
            detected_type or "-",
            status.value,
            confidence,
            len(issues),
        )

        return DocumentVerificationResponse(
            status=status,
            confidence=confidence,
            document_type=expected_type,
            checks=checks,
            issues=issues,
            summary=self._summary(status, issues),
        )

    def _decide(
        self,
        detected_type: Optional[str],
        expected_type: str,
        extracted: ExtractedText,
        name_status: str,
        required_ok: bool,
        cin_ok: bool,
        cin_issue: Optional[str],
        name_issue: Optional[str],
        required_issue: Optional[str],
    ) -> Tuple[VerdictStatus, int, List[str]]:
        """Calcule le statut et la confiance à partir des contrôles."""
        issues: List[str] = []
        confidence = 0

        # Lisibilité du texte (la base de toute vérification).
        if extracted.kind == "image" and not extracted.text.strip():
            confidence += 10
        elif extracted.text.strip():
            confidence += 30
        else:
            confidence += 0

        # Adéquation du type détecté.
        if detected_type is None:
            confidence += 15  # type non confirmé : prudence.
        elif detected_type == expected_type:
            confidence += 35
        else:
            confidence += 5

        # Nom du candidat.
        if name_status == "full":
            confidence += 20
        elif name_status == "partial":
            confidence += 10
        elif name_status == "none":
            confidence += 0

        # Informations requises.
        if required_ok:
            confidence += 15
        else:
            confidence += 5

        # Cohérence exemplaire : bon type + nom complet + infos requises.
        if (
            detected_type == expected_type
            and name_status == "full"
            and required_ok
            and extracted.text.strip()
        ):
            confidence += 10

        # ── Décisions (par ordre de gravité) ──────────────────────────────

        # 1. Type clairement différent du type attendu.
        if detected_type is not None and detected_type != expected_type:
            return (
                VerdictStatus.REJECTED,
                min(confidence, 25),
                [
                    f"Le document fourni ressemble à {DOCUMENT_LABELS[detected_type]} "
                    f"et non à {DOCUMENT_LABELS[expected_type]}."
                ],
            )

        # 2. Aucun texte exploitable (hors image) : rien à vérifier.
        if not extracted.text.strip() and extracted.kind != "image":
            return (
                VerdictStatus.REJECTED,
                min(confidence, 10),
                [
                    "Aucun texte exploitable n'a pu être extrait du document. "
                    "Il ne peut pas être vérifié automatiquement."
                ],
            )

        # 3. Image scannée sans OCR : impossible de confirmer quoi que ce soit.
        if extracted.kind == "image" and not extracted.text.strip():
            return (
                VerdictStatus.REVIEW_REQUIRED,
                min(confidence, 45),
                [
                    "Document image : texte non extractible automatiquement "
                    "(OCR non configuré). Vérification manuelle recommandée."
                ],
            )

        # 4. Nom du candidat absent d'un document lisible : possiblement un
        #    document d'une autre personne.
        if name_status == "none":
            return (
                VerdictStatus.REJECTED,
                max(20, min(confidence, 45)),
                [name_issue or "Le nom du candidat ne figure pas dans le document."],
            )

        # 4 bis. Nom partiellement retrouvé : indice insuffisant, à vérifier.
        if name_status == "partial":
            issues.append(name_issue or "Le nom du candidat n'a pas pu être confirmé dans le document.")
            return (
                VerdictStatus.REVIEW_REQUIRED,
                min(confidence, 84),
                issues,
            )

        # 5. Numéro CIN incohérent (uniquement pour les documents CIN) :
        #    possiblement une carte d'une autre personne — vérification manuelle.
        if expected_type == "CIN" and not cin_ok and cin_issue:
            issues.append(cin_issue)
            return (
                VerdictStatus.REVIEW_REQUIRED,
                min(confidence, 84),
                issues,
            )

        # 6. Informations requises manquantes.
        if not required_ok and required_issue:
            issues.append(required_issue)
            return (
                VerdictStatus.REVIEW_REQUIRED,
                min(confidence, 84),
                issues,
            )

        # 7. Statut final selon le score de confiance (jamais 100 %).
        if confidence >= 85:
            return VerdictStatus.VALID, min(confidence, 95), issues
        if confidence >= 50:
            return VerdictStatus.REVIEW_REQUIRED, confidence, issues
        return VerdictStatus.REJECTED, confidence, issues

    @staticmethod
    def _summary(status: VerdictStatus, issues: List[str]) -> str:
        """Résumé en français, sans jamais affirmer l'authenticité."""
        if status == VerdictStatus.VALID:
            return "Document cohérent avec le dossier du candidat."
        if status == VerdictStatus.REVIEW_REQUIRED:
            if issues:
                return "Vérification recommandée : " + issues[0].rstrip(".") + "."
            return "Document à vérifier manuellement avant validation."
        if issues:
            return "Anomalies détectées : " + issues[0].rstrip(".") + "."
        return "Document rejeté : impossible de le vérifier automatiquement."


# ─────────────────────────────── Construction ────────────────────────────────


def build_document_verification_service(settings=None) -> DocumentVerificationService:
    """Construit le service en fonction de la configuration (.env)."""
    from ..config import get_settings

    settings = settings or get_settings()
    max_bytes = int(getattr(settings, "document_max_bytes", 10 * 1024 * 1024))
    ocr_required = bool(getattr(settings, "document_ocr_required", False))
    return DocumentVerificationService(max_bytes=max_bytes, ocr_required=ocr_required)


# Instance partagée du service (résolue à l'import).
document_verification_service = build_document_verification_service()
